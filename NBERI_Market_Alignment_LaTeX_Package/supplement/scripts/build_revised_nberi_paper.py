from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from copy import deepcopy
import pandas as pd, numpy as np, os, math, shutil

SRC='/mnt/data/manuscript_sources/Completed_News_Based_Exchange_Rate_Risk_and_Volatility_Africa.docx'
OUT='/mnt/data/News_Based_FX_Risk_Market_Alignment_Volatility_Africa_REVISED.docx'
ASSET='/mnt/data/new_paper_assets'

doc=Document(SRC)

# ---------------- styles ----------------
for sname in ['Normal','Body Text']:
    st=doc.styles[sname]
    st.font.name='Times New Roman'; st.font.size=Pt(10.5)
    st.paragraph_format.line_spacing=1.0
    st.paragraph_format.space_after=Pt(3.5 if sname=='Body Text' else 3)
for sname,size in [('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
    if sname in doc.styles:
        st=doc.styles[sname]; st.font.name='Times New Roman'; st.font.size=Pt(size); st.font.bold=True
        st.paragraph_format.line_spacing=1.0
for sec in doc.sections:
    sec.top_margin=Inches(0.7);sec.bottom_margin=Inches(0.68);sec.left_margin=Inches(0.78);sec.right_margin=Inches(0.78)

# Helpers

def set_para_text(p,text,style=None,bold=False,italic=False,align=None,size=None):
    p.clear()
    if style: p.style=style
    r=p.add_run(text);r.bold=bold;r.italic=italic;r.font.name='Times New Roman'
    if size:r.font.size=Pt(size)
    if align is not None:p.alignment=align
    return p

def find_para(starts):
    for p in doc.paragraphs:
        if p.text.strip().startswith(starts): return p
    raise ValueError('Paragraph not found '+starts)

def insert_p_before(anchor,text='',style='Body Text',bold=False,italic=False,align=None):
    new_p=OxmlElement('w:p');anchor._p.addprevious(new_p)
    p=Paragraph(new_p,anchor._parent);p.style=style
    if text:
        r=p.add_run(text);r.bold=bold;r.italic=italic;r.font.name='Times New Roman'
    if align is not None:p.alignment=align
    return p

def insert_table_before(anchor, headers, rows, widths=None, fontsize=8.2):
    table=doc.add_table(rows=1,cols=len(headers));table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=table.rows[0].cells
    for j,h in enumerate(headers):
        hdr[j].text=str(h);hdr[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in hdr[j].paragraphs[0].runs:r.bold=True;r.font.name='Times New Roman';r.font.size=Pt(fontsize)
    for row in rows:
        cells=table.add_row().cells
        for j,val in enumerate(row):
            cells[j].text='' if val is None else str(val)
            cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for par in cells[j].paragraphs:
                par.paragraph_format.space_after=Pt(0);par.paragraph_format.line_spacing=1.0
                for rr in par.runs:rr.font.name='Times New Roman';rr.font.size=Pt(fontsize)
    if widths:
        for row in table.rows:
            for j,w in enumerate(widths):row.cells[j].width=Inches(w)
    tbl=table._tbl;anchor._p.addprevious(tbl)
    return table

def insert_image_before(anchor,path,width=6.5):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path,width=Inches(width))
    anchor._p.addprevious(p._p)
    return p

def delete_para(p):
    el=p._element;el.getparent().remove(el)

# ---------------- title/abstract ----------------
set_para_text(doc.paragraphs[0], 'News-Based Exchange-Rate Risk, Market Alignment, and Volatility Dynamics across African Economies', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=16)
set_para_text(doc.paragraphs[1], 'Continental Panels, GARCH-Family Heterogeneity, Dynamic Transmission, and Out-of-Sample Forecast Evidence', italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)

abstract=(
'This paper evaluates whether a monthly news-based exchange-rate risk measure contains leading information about African currency instability and whether that information becomes stronger in markets where risk narratives are more closely aligned with observed exchange-rate conditions. The source panel covers 54 African countries from January 2015 to December 2025; a pre-specified coverage and return-variation screen yields 51 estimable countries. In the continental benchmark, lagged News-Based Exchange-Rate Risk Index (NBERI) predicts next-month exchange-rate volatility after controlling for volatility history, returns, inflation, country effects, and common time shocks (beta = 0.030, p = 0.004), and remains significant after shared currencies are collapsed to 36 market clusters. The paper then introduces a market-alignment extension based on the 15 eligible countries with the highest contemporaneous correlation between the FX-risk pillar and the bilateral exchange-rate level. In the full 51-country interaction design, the selected group has a total lagged-NBERI slope of 0.063 (p = 0.007), compared with 0.017 (p = 0.147) in the remaining markets; the horizon-1 local-projection slope is also significantly larger in the aligned group. Within the 15-country subsample, NBERI predicts volatility in fixed effects (beta = 0.068, p = 0.013), generates a positive generalized response, explains 3.39 percent of 12-month volatility forecast-error variance, and improves 2023-2025 pooled RMSE by 0.89 percent. Profile-likelihood GARCH-family tests identify FDR-robust conditional-variance effects in Ghana and Libya but not a universal variance channel. The evidence supports NBERI as a leading risk-state indicator whose transmission is economically heterogeneous rather than structurally uniform across African currency markets.'
)
set_para_text(doc.paragraphs[3],abstract,style='Normal')
set_para_text(doc.paragraphs[4],'Keywords: news-based risk; exchange rates; volatility; market alignment; GARCH-X; local projections; Africa',style='Normal')
set_para_text(doc.paragraphs[5],'JEL Classification: F31, C22, C23, C53, G15',style='Normal')

# ---------------- introduction refresh ----------------
p9=find_para('This paper studies the News-Based Exchange-Rate Risk Index')
set_para_text(p9, 'This paper studies the News-Based Exchange-Rate Risk Index (NBERI), a monthly measure of foreign-exchange and external-sector risk narratives. The index is treated strictly as an information-state variable: it is not market volatility, not an exchange-rate transformation, and not a mechanical restatement of realized depreciation. The empirical question is whether the information environment summarized by NBERI precedes subsequent exchange-rate instability after the market\'s own history and conventional macroeconomic controls are taken into account.',style='Body Text')
p10=find_para('The analysis extends the cross-section')
set_para_text(p10, 'The first contribution is continental coverage with explicit market-dependence discipline. The source panel contains 54 countries, of which 51 satisfy pre-model coverage and return-variation criteria. The benchmark therefore uses the broadest defensible cross-section in the supplied monthly data rather than a small hand-picked set of currencies. Currency unions and tightly linked arrangements are then collapsed to 36 effective market clusters in a robustness design, so the main result cannot be attributed to counting repeated CFA- or Common-Monetary-Area exchange-rate paths as independent prices.',style='Body Text')
p11=find_para('The empirical strategy is layered')
set_para_text(p11, 'The second contribution is methodological separation. Two-way fixed effects, PVARs, generalized responses, local projections, and genuine expanding-window forecasts ask whether the risk state contains leading information. Pooled GARCH-X asks a stricter and different question: whether NBERI shifts a common conditional-variance recursion after lagged innovations and variance have absorbed volatility persistence. The paper does not treat agreement across every estimator as a requirement. Instead, it uses disagreement to distinguish predictive information from a universal structural variance parameter.',style='Body Text')
p12=find_para('Three questions organize the paper')
set_para_text(p12, 'The third contribution is a new market-alignment layer designed to organize cross-country heterogeneity without replacing the continental benchmark. Eligible countries are ranked by the contemporaneous correlation between the FX-risk pillar and the country-level bilateral exchange-rate level over 2015-2025; the 15 highest-alignment markets are then examined as a focused heterogeneity subsample. Because exchange-rate levels can be persistent, this ranking is explicitly descriptive rather than causal. Its value is evaluated in the full 51-country panel through selected-group interactions, continuous-alignment interactions, a 1,000-draw random-subset benchmark, horizon-specific local projections, out-of-sample forecasting, and country-level GARCH-family tests. [[FN1]]',style='Body Text')
p13=find_para('The results distinguish a robust leading-state relationship')
set_para_text(p13, 'The results reveal a coherent two-layer pattern. The continental effect is positive, statistically precise, and modest: lagged NBERI predicts next-month volatility with a coefficient of 0.030 (p = 0.004). In the 15 aligned markets the corresponding fixed-effects coefficient rises to 0.068 (p = 0.013), the 12-month NBERI share of volatility forecast-error variance rises to 3.39 percent, and the horizon-1 local-projection effect is 0.121 (p < 0.001). A full-sample interaction test confirms that the aligned group has a larger horizon-1 risk-to-volatility slope (increment = 0.072, p = 0.019). Yet the GARCH-family evidence remains selective: after profile-likelihood refinement and within-family false-discovery-rate correction, robust conditional-variance effects survive only for Ghana under EGARCH-X and Libya under GARCH-X. The paper therefore identifies market alignment as a useful organizing device for heterogeneity while preserving the central conclusion that news-based FX risk is not a universal variance shifter.',style='Body Text')

# ---------------- literature / hypothesis ----------------
p19=find_para('ARCH and GARCH models remain')
set_para_text(p19, 'ARCH and GARCH models remain the standard parametric framework for persistent conditional heteroskedasticity (Engle, 1982; Bollerslev, 1986). EGARCH permits asymmetric responses in log variance (Nelson, 1991), while GJR-GARCH allows negative and positive innovations to affect conditional variance differently (Glosten, Jagannathan, and Runkle, 1993). GARCH-X extensions provide a direct test of whether an observed state variable contributes to conditional variance after endogenous volatility persistence is absorbed (Francq and Thieu, 2019; Pedersen and Rahbek, 2019). In the present setting, that distinction is critical: a news-risk measure can predict next-month instability without necessarily entering every currency\'s variance recursion through the same parameter.',style='Body Text')

h4_anchor=find_para('3. Data and Empirical Design')
insert_p_before(h4_anchor,'H4 - Market-alignment heterogeneity hypothesis. The leading NBERI-volatility relationship should be stronger in markets where the FX-risk pillar is more closely aligned with the observed exchange-rate state, but the alignment screen need not imply a monotonic or universal GARCH-X effect.',style='Body Text')

# ---------------- methodology section 3.9 ----------------
anchor4=find_para('4. Empirical Results')
insert_p_before(anchor4,'3.9 Market-alignment extension and country-level GARCH-family design',style='Heading 2')
insert_p_before(anchor4,'The continental analysis remains the primary specification. The heterogeneity extension ranks the 51 eligible countries by A_i = Corr(FXExternalRisk_it, S_it), where S_it is the bilateral exchange-rate level measured as local-currency units per U.S. dollar over January 2015-December 2025. The 15 highest-alignment markets are Zambia, Ghana, Malawi, Tanzania, Algeria, Burundi, Libya, Nigeria, Tunisia, Morocco, Liberia, Namibia, Congo DRC, Guinea, and Benin. A positive alignment coefficient means that elevated FX-risk narratives tend to occur when the domestic-currency price of the U.S. dollar is high; it is not a statement about next-month depreciation. [[FN2]]',style='Body Text')
insert_p_before(anchor4,'The level-correlation ranking is used only to organize heterogeneity. To guard against a purely mechanical subsample story, the paper estimates the interaction NBERI_i,t-1 x Aligned_i in the full eligible panel, compares the selected-group slope with the remaining-market slope, estimates the same interaction in the horizon-1 local projection, evaluates a continuous NBERI-by-alignment interaction, and benchmarks the selected-group coefficient against 1,000 random 15-country subsets. Spearman and detrended log-exchange-rate correlations are reported as descriptive sensitivity checks. The interaction tests, rather than the raw ranking itself, provide the formal evidence on differential transmission.',style='Body Text')
insert_p_before(anchor4,'Within the aligned subsample, the paper reruns the fixed-effects, PVAR, generalized-response, FEVD, local-projection, forecast, absolute-return, and first-difference-NBERI exercises. Country-level conditional variance is then examined using six models: GARCH(1,1), GARCH-X, EGARCH(1,1), EGARCH-X, GJR-GARCH/TGARCH(1,1), and GJR-GARCH/TGARCH-X. Each model uses an AR(1) conditional mean and standardized Student-t innovations. The X models include lagged standardized NBERI in the variance equation. Likelihood-ratio tests compare each X model with its nested non-X family, and Benjamini-Hochberg q-values control multiplicity across the 15 countries within each family.',style='Body Text')
insert_p_before(anchor4,'Because nonlinear volatility likelihoods can be sensitive to starting values, the reported country-level X estimates are subjected to an explicit profile-likelihood check over the NBERI variance coefficient before joint refinement. The profile step ensures that a reported X effect improves the corresponding nested family rather than reflecting a single local optimum. Residual squared-autocorrelation and ARCH-LM diagnostics are reported alongside the likelihood tests. [[FN3]]',style='Body Text')

# ---------------- new empirical sections before Discussion ----------------
anchor_disc=find_para('5. Discussion')
insert_p_before(anchor_disc,'4.9 FX-risk/exchange-rate alignment across the top 15 markets',style='Heading 2')
insert_p_before(anchor_disc,'Table 8. FX-risk pillar and exchange-rate-level alignment, 2015-2025',style='Caption')
align=pd.read_csv('/mnt/data/nberi_alignment_interaction/selected15_level_alignment_robustness.csv')
order=['Zambia','Ghana','Malawi','Tanzania','Algeria','Burundi','Libya','Nigeria','Tunisia','Morocco','Liberia','Namibia','Congo DRC','Guinea','Benin']
align=align.set_index('Country').reindex(order).reset_index()
rows=[]
for _,r in align.iterrows():
    rows.append([r.Country,f'{r.Raw_Pearson_level:.3f}',f'{r.Spearman_level:.3f}',f'{r.Linear_detrended_loglevel_corr:.3f}',f'{r.Cubic_detrended_loglevel_corr:.3f}'])
insert_table_before(anchor_disc,['Country','Raw Pearson','Spearman','Linear-detrended log S','Cubic-detrended log S'],rows,widths=[1.35,.8,.8,1.15,1.15],fontsize=7.8)
insert_p_before(anchor_disc,'Notes: Raw Pearson correlation is the ranking statistic used to define the top-15 alignment group. Detrended columns correlate residuals after removing country-specific linear or cubic time trends from log exchange rates and the FX-risk pillar. These sensitivity columns diagnose whether raw level co-movement is dominated by common trending behavior; they are not used to re-select countries.',style='Normal')
insert_image_before(anchor_disc,os.path.join(ASSET,'figure8_alignment.png'),width=6.3)
insert_p_before(anchor_disc,'Figure 8. Top 15 FX-risk/exchange-rate aligned markets.',style='Caption')
insert_p_before(anchor_disc,'The alignment ranking is economically heterogeneous rather than a simple proxy for trend. Raw correlations range from 0.718 in Zambia to 0.247 in Benin. Rank correlations remain positive for all 15 markets and exceed 0.40 in several cases, while detrended correlations are strong in Algeria, Nigeria, Namibia, Liberia, Ghana, Malawi, and Libya but much weaker in Zambia, Tanzania, and Guinea. This pattern is precisely why the level correlation is treated as a screening statistic rather than an identifying variable: it identifies markets in which the news-risk state and the observed FX state co-move strongly in levels, but the mechanism behind that co-movement differs across countries.',style='Body Text')

insert_p_before(anchor_disc,'4.10 Does alignment identify stronger risk-to-volatility transmission?',style='Heading 2')
insert_p_before(anchor_disc,'Table 9. Full-sample tests of alignment heterogeneity',style='Caption')
feint=pd.read_csv('/mnt/data/nberi_alignment_interaction/group_interaction_FE.csv')
lpint=pd.read_csv('/mnt/data/nberi_alignment_interaction/group_interaction_LP_h1.csv')
cont=pd.read_csv('/mnt/data/nberi_alignment_interaction/continuous_alignment_FE.csv')
perm=pd.read_csv('/mnt/data/nberi_alignment_interaction/random_subset_permutation_summary.csv')
rows=[
['FE: remaining eligible slope',f"{feint.iloc[0].Estimate:.3f}",f"{feint.iloc[0].SE:.3f}",f"{feint.iloc[0].p:.3f}"],
['FE: selected-15 incremental slope',f"{feint.iloc[1].Estimate:.3f}",f"{feint.iloc[1].SE:.3f}",f"{feint.iloc[1].p:.3f}"],
['FE: selected-15 total slope',f"{feint.iloc[2].Estimate:.3f}",f"{feint.iloc[2].SE:.3f}",f"{feint.iloc[2].p:.3f}"],
['LP h=1: remaining eligible slope',f"{lpint.iloc[0].Estimate:.3f}",f"{lpint.iloc[0].SE:.3f}",f"{lpint.iloc[0].p:.3f}"],
['LP h=1: selected-15 incremental slope',f"{lpint.iloc[1].Estimate:.3f}",f"{lpint.iloc[1].SE:.3f}",f"{lpint.iloc[1].p:.3f}"],
['LP h=1: selected-15 total slope',f"{lpint.iloc[2].Estimate:.3f}",f"{lpint.iloc[2].SE:.3f}",f"{lpint.iloc[2].p:.3f}"],
['Continuous NBERI x alignment interaction',f"{cont.iloc[1].Estimate:.3f}",f"{cont.iloc[1].SE:.3f}",f"{cont.iloc[1].p:.3f}"],
['Random-subset benchmark: selected total slope',f"{perm.iloc[1].Observed:.3f}",'94.4th percentile',f"{perm.iloc[1].Empirical_one_sided_p:.3f}"],
]
insert_table_before(anchor_disc,['Test','Estimate','SE / benchmark','p-value'],rows,widths=[3.1,.75,1.1,.7],fontsize=8.0)
insert_p_before(anchor_disc,'Notes: Fixed-effects and local-projection interaction models are estimated in the full eligible-country panel with country and calendar-time effects and the same dynamic controls used in the benchmark. The random-subset benchmark compares the selected group with 1,000 draws of 15 eligible countries and reports an empirical one-sided probability.',style='Normal')
insert_image_before(anchor_disc,os.path.join(ASSET,'figure9_group_interaction.png'),width=6.2)
insert_p_before(anchor_disc,'Figure 9. Risk-to-volatility slopes in aligned and remaining eligible markets.',style='Caption')
insert_p_before(anchor_disc,'The group comparison provides the formal evidence that the alignment screen is associated with stronger transmission. In the remaining eligible markets, the lagged-NBERI coefficient is 0.017 (p = 0.147). The selected-15 total slope is 0.063 (p = 0.007), with an incremental selected-group effect of 0.046 (p = 0.074). The horizon-1 result is sharper: the remaining-market slope is 0.029 (p = 0.032), the selected-group slope is 0.100 (p < 0.001), and the selected-group increment is 0.072 (p = 0.019). The selected-group fixed-effects slope lies near the 94th percentile of 1,000 random 15-country subsets, with a one-sided empirical probability of approximately 0.057. This is suggestive rather than decisive permutation evidence, but it shows that the top-15 result is not typical of an arbitrary 15-country draw.',style='Body Text')
insert_p_before(anchor_disc,'The continuous interaction between NBERI and the country-level alignment coefficient is positive but insignificant (0.012, p = 0.412). Alignment therefore does not act as a linear dose-response moderator. The evidence is better described as a grouped market-state distinction: the highest-alignment markets collectively display stronger short-horizon transmission, but a larger raw level correlation does not mechanically imply a proportionally larger volatility coefficient.',style='Body Text')

insert_p_before(anchor_disc,'4.11 Dynamic and forecast evidence in the aligned-market subsample',style='Heading 2')
insert_p_before(anchor_disc,'Table 10. Dynamic and forecast evidence in the 15 aligned markets',style='Caption')
fe15=pd.read_csv('/mnt/data/nberi_15_country_full_extension/03_two_way_FE.csv')
pv15=pd.read_csv('/mnt/data/nberi_15_country_full_extension/06_PVAR_coefficients.csv')
irf15=pd.read_csv('/mnt/data/nberi_15_country_full_extension/07_generalized_IRF.csv')
fevd15=pd.read_csv('/mnt/data/nberi_15_country_full_extension/08_generalized_FEVD_12m.csv')
lp15=pd.read_csv('/mnt/data/nberi_15_country_full_extension/09_local_projections.csv')
fs15=pd.read_csv('/mnt/data/nberi_15_country_full_extension/11_forecast_pooled_summary.csv')
rb15=pd.read_csv('/mnt/data/nberi_15_country_full_extension/13_robustness.csv')
fe_row=fe15[(fe15.Specification=='Baseline + inflation')&(fe15.Regressor=='L_nberi_z')].iloc[0]
pv_row=pv15[(pv15.Equation=='vol_z')&(pv15.Lagged_regressor=='L_nberi_z')].iloc[0]
irf1=irf15[irf15.Horizon==1].iloc[0];lpr1=lp15[lp15.Horizon==1].iloc[0]
rmse_imp=fs15.loc[fs15.Metric=='RMSE + NBERI','Improvement_pct'].iloc[0];ql_imp=fs15.loc[fs15.Metric=='QLIKE + NBERI','Improvement_pct'].iloc[0];cw=fs15.loc[fs15.Metric=='Clark-West mean'].iloc[0];qd=fs15.loc[fs15.Metric=='QLIKE loss difference mean'].iloc[0]
absrow=rb15[rb15.Specification=='Absolute-return outcome'].iloc[0];diffrow=rb15[rb15.Specification=='First-difference NBERI'].iloc[0]
rows=[
['Two-way FE: lagged NBERI',f'{fe_row.Coefficient:.3f}',f'{fe_row.DK_SE:.3f}',f'{fe_row.p_value:.3f}'],
['PVAR volatility equation: lagged NBERI',f'{pv_row.Coefficient:.3f}',f'{pv_row.DK_SE:.3f}',f'{pv_row.p_value:.3f}'],
['Generalized IRF, h=1',f'{irf1.FX_volatility_response:.3f}',f'95% CI [{irf1.CI95_low:.3f}, {irf1.CI95_high:.3f}]',''],
['12-month FEVD share from NBERI',f"{100*fevd15.loc[fevd15.Shock=='NBERI','Share_12m_FX_volatility_FEVD'].iloc[0]:.2f}%",'',''],
['Local projection, h=1',f'{lpr1.NBERI_beta:.3f}',f'{lpr1.DK_SE:.3f}',f'{lpr1.p_value:.4f}'],
['Absolute-return robustness',f'{absrow.Coefficient:.3f}',f'{absrow.DK_SE:.3f}',f'{absrow.p_value:.3f}'],
['First-difference NBERI robustness',f'{diffrow.Coefficient:.3f}',f'{diffrow.DK_SE:.3f}',f'{diffrow.p_value:.3f}'],
['Forecast RMSE improvement',f'{rmse_imp:.2f}%','',''],
['Forecast QLIKE improvement',f'{ql_imp:.2f}%','',''],
['Clark-West nested test',f'{cw.Value:.3f}','','p = '+f'{cw.p_value:.4f}'],
['QLIKE loss-difference test',f'{qd.Value:.3f}','','p = '+f'{qd.p_value:.3f}'],
]
insert_table_before(anchor_disc,['Result','Estimate','SE / interval','p-value'],rows,widths=[2.65,.9,1.55,.8],fontsize=8.0)
insert_image_before(anchor_disc,os.path.join(ASSET,'figure10_lp15.png'),width=6.2)
insert_p_before(anchor_disc,'Figure 10. Local-projection response in the 15 aligned markets.',style='Caption')
insert_p_before(anchor_disc,'The aligned-market subsample strengthens the dynamic signal without changing its interpretation. The fixed-effects coefficient is 0.068 (p = 0.013), more than twice the continental coefficient. The PVAR returns the same one-step coefficient and remains dynamically stable, with a maximum companion-root modulus of 0.483. A one-standard-deviation NBERI innovation generates a generalized volatility response of 0.099 at horizon 1, with a 95 percent country-block bootstrap interval of 0.054 to 0.131, and 0.056 at horizon 2. NBERI accounts for 3.39 percent of the 12-month volatility forecast-error variance, compared with 1.05 percent in the continental system.',style='Body Text')
insert_p_before(anchor_disc,'Local projections show a broad near-term response: the h=1 coefficient is 0.121 (p < 0.001), h=2 is 0.130 (p < 0.001), and the response remains positive through several subsequent horizons. The absolute-return specification is also positive (0.066, p = 0.023), whereas the first difference of NBERI is insignificant (-0.025, p = 0.406). As in the continental results, the information appears to reside in a persistent risk state rather than a one-month jump in the index.',style='Body Text')

insert_p_before(anchor_disc,'4.12 Country-level GARCH-family evidence and out-of-sample gains',style='Heading 2')
insert_p_before(anchor_disc,'Table 11. Profile-likelihood GARCH-family results surviving within-family FDR correction',style='Caption')
prof=pd.read_csv('/mnt/data/nberi_15_country_garch_profiled/profiled_sig_q10.csv')
rows=[]
for _,r in prof.iterrows():
    rows.append([r.Country,r.Family+'-X',f'{r.delta:.3f}',f'{r.LR_p:.4f}',f'{r.BH_FDR_q:.4f}',f'{r.LB_sq5_p:.3f}',f'{r.ARCH_LM5_p:.3f}'])
insert_table_before(anchor_disc,['Country','Model','delta','LR p','BH q','LB(5) sq. p','ARCH-LM(5) p'],rows,widths=[1.0,1.05,.65,.65,.65,.8,.9],fontsize=7.8)
insert_p_before(anchor_disc,'Notes: Reported rows are the only country-family X effects that survive Benjamini-Hochberg q < 0.10 after profile-likelihood refinement of the NBERI variance coefficient. Forty-five X models are tested in total: GARCH-X, EGARCH-X, and GJR/TGARCH-X for each of 15 countries. Full results are reported in Appendix H.',style='Normal')
insert_image_before(anchor_disc,os.path.join(ASSET,'figure11_garch_profiled.png'),width=6.4)
insert_p_before(anchor_disc,'Figure 11. Profile-likelihood NBERI variance coefficients across the 15-country GARCH family.',style='Caption')
insert_p_before(anchor_disc,'The profile-likelihood exercise sharply disciplines the country-level GARCH evidence. Only two effects survive within-family false-discovery control: Ghana under EGARCH-X (delta = 0.766; LR p = 0.00024; q = 0.0036) and Libya under GARCH-X (delta = 0.645; LR p = 0.00050; q = 0.0074). Libya passes the squared-residual and ARCH diagnostics comfortably. Ghana\'s ARCH-LM diagnostic is acceptable at conventional levels, although the Ljung-Box test on squared standardized residuals is borderline at p = 0.044; the Ghana result is therefore interpreted with additional caution. The remaining 43 country-family X tests do not survive q < 0.10 after the profile-likelihood refinement. This result is more informative than a large collection of nominal country p-values: it shows that conditional-variance transmission exists in specific markets but is not a general feature of the aligned sample.',style='Body Text')

insert_image_before(anchor_disc,os.path.join(ASSET,'figure12_forecast15.png'),width=6.25)
insert_p_before(anchor_disc,'Figure 12. Country-level QLIKE improvement from adding NBERI in the aligned-market forecast experiment.',style='Caption')
insert_p_before(anchor_disc,'The forecast horse race adds a genuine prediction test. Across the aligned markets, adding lagged NBERI lowers pooled RMSE by 0.89 percent, from 2.254 to 2.234, and the Clark-West test favors the augmented model (p = 0.0005). Mean QLIKE falls by 2.85 percent, although the HAC test of the pooled QLIKE loss difference is not significant at conventional levels (p = 0.123). The country distribution is heterogeneous: Ghana records the largest QLIKE improvement, followed by Morocco and Burundi, while several countries show small deteriorations. The forecast evidence therefore supports incremental predictive content without implying uniform country-level dominance.',style='Body Text')

# ---------------- Discussion rewrite ----------------
p108=find_para('5.1 A leading risk state')
set_para_text(p108,'5.1 Continental leading information and the boundary of the claim',style='Heading 2')
p109=find_para('The evidence is most coherent')
set_para_text(p109,'The continental evidence establishes the baseline result: NBERI contains leading information about subsequent monthly exchange-rate instability. The coefficient survives market-history controls, inflation, country effects, common time shocks, and the collapse of repeated currency-union paths. PVARs, generalized responses, local projections, absolute-return robustness, and out-of-sample tests all point in the same direction. The size is economically modest, which is appropriate for an incremental information variable rather than a stand-alone market model.',style='Body Text')
p110=find_para('The pooled GARCH-X result')
set_para_text(p110,'The pooled GARCH-X null remains an important boundary. The common monthly variance recursion is extremely persistent and does not admit a statistically meaningful continent-wide NBERI variance coefficient. The new country-family exercise does not overturn that result; it explains it. After profile-likelihood refinement and multiplicity control, only Ghana and Libya exhibit robust family-specific X effects. The absence of a universal variance parameter is therefore not evidence that NBERI lacks information. It indicates that the information enters market dynamics through heterogeneous channels rather than one common GARCH mechanism.',style='Body Text')
p111=find_para('5.2 Cross-market heterogeneity')
set_para_text(p111,'5.2 Market alignment as an organizing device for heterogeneity',style='Heading 2')
p112=find_para('The expanded sample makes it possible')
set_para_text(p112,'The market-alignment extension provides the paper\'s main heterogeneity result. The top-15 aligned markets have a larger total fixed-effects slope than the remaining eligible countries and a significantly larger horizon-1 local-projection response. The selected-group coefficient also lies near the 94th percentile of random 15-country subsets. At the same time, the continuous alignment interaction is insignificant, so the raw correlation should not be interpreted as a structural moderator measured on a cardinal scale. The useful insight is categorical and empirical: markets in which FX-risk narratives and the observed exchange-rate state are most tightly synchronized, as a group, transmit the news-risk state into subsequent volatility more strongly.',style='Body Text')
p113=find_para('5.3 Surveillance and policy interpretation')
set_para_text(p113,'5.3 Why the GARCH-family and forecast results matter',style='Heading 2')
p114=find_para('For macro-financial surveillance')
set_para_text(p114,'The combined evidence separates three objects that are often conflated: information, conditional variance, and forecasting value. NBERI is informative in the panel and local-projection sense; conditional-variance effects are concentrated in a small number of countries; and out-of-sample gains are positive on average but uneven across markets and loss functions. This distinction is operationally useful. A surveillance system can treat NBERI as a state variable that enriches the information set without assuming that it enters every currency\'s variance equation or improves every forecast origin.',style='Body Text')
# add 5.4 before conclusion
anchor_conc=find_para('6. Conclusion')
insert_p_before(anchor_conc,'5.4 Contribution to empirical FX-risk measurement',style='Heading 2')
insert_p_before(anchor_conc,'The paper contributes an empirical architecture for evaluating text-based market-risk measures rather than merely documenting correlation. The architecture begins with a broad panel, separates news risk from market outcomes, tests effective currency-market dependence, distinguishes predictive regressions from conditional-variance recursions, traces dynamic propagation, imposes real-time forecast evaluation, and then uses an alignment screen only as a disciplined heterogeneity extension. The resulting conclusion is stronger because it is narrower: NBERI is a leading information state with economically meaningful market heterogeneity, not a universal structural volatility factor.',style='Body Text')

# ---------------- Conclusion rewrite ----------------
p116=find_para('This paper evaluates a news-based exchange-rate risk measure')
set_para_text(p116,'This paper evaluates a news-based exchange-rate risk measure across the broadest African sample supported by the supplied monthly data and adds a new market-alignment framework for understanding why the signal is stronger in some currency markets than others. Of 54 source countries, 51 pass pre-specified coverage and variation screens. The design keeps the news measure separate from market outcomes, adjusts inference for serial and cross-sectional dependence, collapses shared-currency paths, and combines panel prediction, dynamic propagation, conditional-variance models, and real-time forecast evaluation in one framework.',style='Body Text')
p117=find_para('The central result is a positive but moderate leading relationship')
set_para_text(p117,'The continental result is a positive but moderate leading relationship: a one-standard-deviation increase in lagged NBERI predicts approximately 0.03 standard deviations more next-month FX volatility after controls, and the relationship survives the 36-cluster currency-market robustness. The alignment extension adds a second result. In the top 15 FX-risk/exchange-rate aligned markets, the fixed-effects coefficient is 0.068 (p = 0.013), the horizon-1 local-projection response is 0.121 (p < 0.001), and NBERI explains 3.39 percent of the 12-month volatility forecast-error variance. Full-sample interaction tests show that the aligned group has a significantly larger horizon-1 slope than the remaining eligible markets.',style='Body Text')
p118=find_para('Out-of-sample evidence is stronger')
set_para_text(p118,'Out-of-sample performance is consistent with incremental rather than dominant predictive value. In the aligned markets, NBERI lowers pooled RMSE by 0.89 percent and passes the Clark-West nested forecast test, while the pooled QLIKE improvement is positive but not statistically decisive under the HAC loss-difference test. The country-level GARCH family provides an equally important boundary: after profile-likelihood refinement and false-discovery correction, robust NBERI variance effects survive only in Ghana and Libya. The common pooled GARCH-X null and the selective country-family evidence are therefore mutually consistent.',style='Body Text')
p119=find_para('The empirical contribution lies in this separation')
set_para_text(p119,'The paper\'s novel contribution is the separation of broad informational content from market-specific volatility transmission. News-based FX risk contains measurable leading information across the continent, but the strength and mechanism of that information depend on market context. The market-alignment layer identifies a group with stronger short-horizon transmission without converting a descriptive level correlation into a causal claim. For surveillance, this implies that text-based FX-risk measures are most useful as disciplined additions to the market information set, with country-specific model validation determining when they should enter conditional-variance or forecasting systems.',style='Body Text')

# ---------------- Limitations shorten / update ----------------
p121=find_para('The analysis is constrained by the frequency')
set_para_text(p121,'The main limitation is data frequency. Monthly exchange rates cannot recover intramonth quadratic variation, so the panel outcome is a monthly volatility proxy and the GARCH-family models have relatively short time series. This is why the panel, dynamic, and forecast evidence carries more weight than any single country-level variance recursion. Higher-frequency exchange-rate data would permit sharper GARCH-X, EGARCH-X, and threshold-volatility tests.',style='Body Text')
p122=find_para('African exchange-rate observations are also not cross-sectionally independent')
set_para_text(p122,'A second limitation concerns the market-alignment screen. Exchange-rate levels can be persistent or trending, and a contemporaneous level correlation is not an identification strategy. The paper therefore treats the top-15 ranking as descriptive, reports rank and detrended sensitivity measures, performs full-sample interaction tests, and compares the selected group with random 15-country subsets. These checks strengthen the heterogeneity interpretation but do not make the alignment coefficient causal.',style='Body Text')
p123=find_para('NBERI is persistent and observational')
set_para_text(p123,'Finally, NBERI is observational and the 2023-2025 forecast evaluation window is short relative to the full estimation period. Lagging the index, fixed effects, reverse PVAR paths, state-versus-change robustness, profile-likelihood volatility tests, multiplicity control, and genuine out-of-sample evaluation narrow several alternative explanations but do not identify an exogenous structural shock. The paper therefore maintains predictive language and treats country-specific mechanisms as hypotheses for future work with higher-frequency prices and richer institutional covariates.',style='Body Text')
p124=find_para('Finally, the out-of-sample window covers')
delete_para(p124)

# ---------------- References: insert Nelson ----------------
ref_anchor=find_para('Olowe, R. A.')
insert_p_before(ref_anchor,'Nelson, D. B. (1991). Conditional heteroskedasticity in asset returns: A new approach. Econometrica, 59(2), 347-370.',style='Normal')

# ---------------- Appendices G-I ----------------
# Append at end
p=doc.add_paragraph('Appendix G. Market-Alignment Robustness',style='Heading 1')
doc.add_paragraph('The market-alignment ranking is descriptive and is not used as a causal identifying assumption. Table G1 reports rank-based and detrended sensitivity measures for the selected countries. Positive raw correlation is common across the group, but detrended associations vary materially, which is consistent with using alignment as a heterogeneity screen rather than a structural parameter.',style='Body Text')
doc.add_paragraph('Table G1. Market-alignment sensitivity diagnostics',style='Caption')
t=doc.add_table(rows=1,cols=5);t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
headers=['Country','Raw Pearson','Spearman','Linear detrended','Cubic detrended']
for j,h in enumerate(headers):t.rows[0].cells[j].text=h
for _,r in align.iterrows():
    cells=t.add_row().cells;vals=[r.Country,f'{r.Raw_Pearson_level:.3f}',f'{r.Spearman_level:.3f}',f'{r.Linear_detrended_loglevel_corr:.3f}',f'{r.Cubic_detrended_loglevel_corr:.3f}']
    for j,v in enumerate(vals):cells[j].text=v
for row in t.rows:
    for cell in row.cells:
        for pa in cell.paragraphs:
            for rr in pa.runs:rr.font.name='Times New Roman';rr.font.size=Pt(7.8)

hpar=doc.add_paragraph('Appendix H. Full Profile-Likelihood GARCH-Family Results',style='Heading 1')
hpar.paragraph_format.page_break_before=True
doc.add_paragraph('Tables H1a-H1c report all 45 country-family X tests after profile-likelihood refinement. BH q-values are computed separately within GARCH-X, EGARCH-X, and GJR/TGARCH-X across the 15 countries.',style='Body Text')
allg=pd.read_csv('/mnt/data/nberi_15_country_garch_profiled/profiled_garch_x_results.csv')
for fam,label in [('GARCH','Table H1a. GARCH-X country results'),('EGARCH','Table H1b. EGARCH-X country results'),('GJR/TGARCH','Table H1c. GJR/TGARCH-X country results')]:
    doc.add_paragraph(label,style='Caption')
    t=doc.add_table(rows=1,cols=8);t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    headers=['Country','delta','LR p','BH q','AIC','LB sq. p','ARCH-LM p','Profile grid']
    for j,h in enumerate(headers):t.rows[0].cells[j].text=h
    sub=allg[allg.Family==fam].copy().sort_values('Country')
    for _,r in sub.iterrows():
        cells=t.add_row().cells;vals=[r.Country,f'{r.delta:.3f}',f'{r.LR_p:.4f}',f'{r.BH_FDR_q:.4f}',f'{r.AIC:.1f}',f'{r.LB_sq5_p:.3f}',f'{r.ARCH_LM5_p:.3f}',f'{r.profile_grid_best_delta:.2f}']
        for j,v in enumerate(vals):cells[j].text=v
    for row in t.rows:
        for cell in row.cells:
            for pa in cell.paragraphs:
                pa.paragraph_format.space_after=Pt(0)
                for rr in pa.runs:rr.font.name='Times New Roman';rr.font.size=Pt(7.6)

ipar=doc.add_paragraph('Appendix I. Random-Subset Benchmark',style='Heading 1')
ipar.paragraph_format.page_break_before=True
doc.add_paragraph('To assess whether the selected-15 fixed-effects slope is unusually large relative to arbitrary country groupings, 1,000 subsets of 15 countries are drawn without replacement from the 51 eligible markets. The selected group\'s incremental slope is at the 94.1st percentile of the random distribution and its total slope is at the 94.4th percentile. The corresponding one-sided empirical probabilities are 0.060 and 0.057. These values are suggestive rather than conventionally significant and are interpreted accordingly.',style='Body Text')
doc.add_paragraph('Table I1. Random 15-country subset benchmark',style='Caption')
perm2=pd.read_csv('/mnt/data/nberi_alignment_interaction/random_subset_permutation_summary.csv')
t=doc.add_table(rows=1,cols=6);t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
headers=['Statistic','Observed','Random mean','Random SD','Percentile','Empirical p']
for j,h in enumerate(headers):t.rows[0].cells[j].text=h
for _,r in perm2.iterrows():
    cells=t.add_row().cells;vals=[r.Statistic,f'{r.Observed:.3f}',f'{r.Random_mean:.3f}',f'{r.Random_sd:.3f}',f'{r.Percentile:.1f}',f'{r.Empirical_one_sided_p:.3f}']
    for j,v in enumerate(vals):cells[j].text=v
for row in t.rows:
    for cell in row.cells:
        for pa in cell.paragraphs:
            for rr in pa.runs:rr.font.name='Times New Roman';rr.font.size=Pt(8)

# Table header bold and repeat header for all tables
for table in doc.tables:
    if table.rows:
        trPr=table.rows[0]._tr.get_or_add_trPr();tblHeader=OxmlElement('w:tblHeader');tblHeader.set(qn('w:val'),'true');trPr.append(tblHeader)
        for cell in table.rows[0].cells:
            for rr in cell.paragraphs[0].runs:rr.bold=True
    table.alignment=WD_TABLE_ALIGNMENT.CENTER

# Global font cleanup for added paragraphs and ensure compact spacing
for p in doc.paragraphs:
    if p.style.name in ['Normal','Body Text']:
        p.paragraph_format.line_spacing=1.0
        if p.paragraph_format.space_after is None:p.paragraph_format.space_after=Pt(3)
    for r in p.runs:
        if not r.font.name:r.font.name='Times New Roman'

# Save
if os.path.exists(OUT):os.remove(OUT)
doc.save(OUT)
print(OUT)
